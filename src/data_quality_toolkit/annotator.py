from pathlib import Path
from shutil import copy2

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from src.data_quality_toolkit.utils import source_row_number
from src.data_quality_toolkit.validator import ValidationResult


MISSING_VALUE_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFFF99",
)

INVALID_VALUE_FILL = PatternFill(
    fill_type="solid",
    fgColor="FF9999",
)

DUPLICATE_ROW_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFD699",
)

UNEXPECTED_COLUMN_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9D9D9",
)


def annotate_workbook(
    input_path: str | Path,
    result: ValidationResult,
    output_path: str | Path,
) -> None:
    """
    Create a reviewed copy of an Excel workbook with validation
    issues highlighted.

    Args:
        input_path: Existing XLSX workbook.
        result: Validation issues detected for the workbook data.
        output_path: Destination for the annotated workbook.
    """
    source = Path(input_path)
    destination = Path(output_path)

    if not source.exists():
        raise FileNotFoundError(
            f"Input workbook not found: {source}"
        )

    if source.suffix.lower() != ".xlsx":
        raise ValueError(
            "Annotated workbook output currently requires "
            "an .xlsx input file."
        )

    destination.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    copy2(
        source,
        destination,
    )

    workbook = load_workbook(destination)
    worksheet = workbook.active

    header_columns = {
        cell.value: cell.column
        for cell in worksheet[1]
        if cell.value is not None
    }

    for column in result.unexpected_columns:
        column_number = header_columns.get(column)

        if column_number is not None:
            worksheet.cell(
                row=1,
                column=column_number,
            ).fill = UNEXPECTED_COLUMN_FILL

    for column, rows in result.missing_values.items():
        column_number = header_columns.get(column)

        if column_number is None:
            continue

        for row_index in rows:
            worksheet.cell(
                row=source_row_number(row_index),
                column=column_number,
            ).fill = MISSING_VALUE_FILL

    for column, rows in result.invalid_values.items():
        column_number = header_columns.get(column)

        if column_number is None:
            continue

        for row_index in rows:
            worksheet.cell(
                row=source_row_number(row_index),
                column=column_number,
            ).fill = INVALID_VALUE_FILL

    # 1. Duplicate rows first
    for row_index in result.duplicate_rows:
        excel_row = source_row_number(row_index)

        for column_number in range(
                1,
                worksheet.max_column + 1,
        ):
            worksheet.cell(
                row=excel_row,
                column=column_number,
            ).fill = DUPLICATE_ROW_FILL

    # 2. Missing values second
    for column, rows in result.missing_values.items():
        column_number = header_columns.get(column)

        if column_number is None:
            continue

        for row_index in rows:
            worksheet.cell(
                row=source_row_number(row_index),
                column=column_number,
            ).fill = MISSING_VALUE_FILL

    # 3. Invalid values last
    for column, rows in result.invalid_values.items():
        column_number = header_columns.get(column)

        if column_number is None:
            continue

        for row_index in rows:
            worksheet.cell(
                row=source_row_number(row_index),
                column=column_number,
            ).fill = INVALID_VALUE_FILL

    add_legend_sheet(workbook)

    workbook.save(destination)

def add_legend_sheet(workbook) -> None:
    """
    Add a legend explaining the workbook highlight colors.
    """
    if "Legend" in workbook.sheetnames:
        del workbook["Legend"]

    legend = workbook.create_sheet("Legend")

    legend.append(
        [
            "Highlight",
            "Meaning",
        ]
    )

    legend.append(
        [
            "",
            "Missing required value",
        ]
    )

    legend.append(
        [
            "",
            "Invalid value",
        ]
    )

    legend.append(
        [
            "",
            "Duplicate row",
        ]
    )

    legend.append(
        [
            "",
            "Unexpected column",
        ]
    )

    legend["A2"].fill = MISSING_VALUE_FILL
    legend["A3"].fill = INVALID_VALUE_FILL
    legend["A4"].fill = DUPLICATE_ROW_FILL
    legend["A5"].fill = UNEXPECTED_COLUMN_FILL

    legend["A1"].font = Font(bold=True)
    legend["B1"].font = Font(bold=True)

    legend.column_dimensions["A"].width = 14
    legend.column_dimensions["B"].width = 28