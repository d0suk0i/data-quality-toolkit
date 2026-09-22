import pandas as pd

from src.data_quality_toolkit.exporter import export_report_csv
from src.data_quality_toolkit.validator import ValidationResult
from openpyxl import load_workbook


def test_export_report_csv_creates_report(tmp_path):
    result = ValidationResult(
        missing_columns=["email"],
        unexpected_columns=["notes"],
        missing_values={"name": [2]},
        duplicate_rows=[4],
        invalid_values={"employment_type": [1, 3]},
    )

    output_path = tmp_path / "report.csv"

    export_report_csv(result, output_path)

    report = pd.read_csv(output_path)

    assert list(report.columns) == [
        "issue_type",
        "column",
        "source_row",
        "details",
    ]

    assert len(report) == 6

    assert "missing_column" in report["issue_type"].values
    assert "unexpected_column" in report["issue_type"].values
    assert "missing_value" in report["issue_type"].values
    assert "duplicate_row" in report["issue_type"].values
    assert "invalid_value" in report["issue_type"].values


def test_export_report_csv_handles_valid_result(tmp_path):
    result = ValidationResult()

    output_path = tmp_path / "report.csv"

    export_report_csv(result, output_path)

    report = pd.read_csv(output_path)

    assert len(report) == 0
    assert list(report.columns) == [
        "issue_type",
        "column",
        "source_row",
        "details",
    ]

def test_export_report_excel_creates_summary_and_issues_sheets(tmp_path):
    from src.data_quality_toolkit.exporter import export_report_excel

    result = ValidationResult(
        missing_columns=["email"],
        missing_values={"name": [2]},
        duplicate_rows=[4],
        invalid_values={"employment_type": [1]},
    )

    output_path = tmp_path / "report.xlsx"

    export_report_excel(result, output_path)

    workbook = load_workbook(output_path)

    assert "Summary" in workbook.sheetnames
    assert "Issues" in workbook.sheetnames

    summary = workbook["Summary"]
    issues = workbook["Issues"]

    assert summary["A1"].value == "Validation Status"
    assert summary["B1"].value == "FAILED"

    assert issues["A1"].value == "issue_type"
    assert issues["B1"].value == "column"
    assert issues["C1"].value == "source_row"
    assert issues["D1"].value == "details"


def test_export_report_excel_handles_valid_result(tmp_path):
    from src.data_quality_toolkit.exporter import export_report_excel

    result = ValidationResult()

    output_path = tmp_path / "report.xlsx"

    export_report_excel(result, output_path)

    workbook = load_workbook(output_path)

    summary = workbook["Summary"]
    issues = workbook["Issues"]

    assert summary["B1"].value == "PASSED"
    assert issues.max_row == 1