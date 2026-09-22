import pandas as pd

from src.data_quality_toolkit.cli import main


def test_cli_returns_success_for_valid_file(tmp_path, capsys):
    data = pd.DataFrame(
        {
            "name": ["Alice", "Bob"],
            "email": [
                "alice@example.com",
                "bob@example.com",
            ],
        }
    )

    file_path = tmp_path / "valid.csv"
    data.to_csv(file_path, index=False)

    exit_code = main(
        [
            str(file_path),
            "--required",
            "name",
            "email",
        ]
    )

    output = capsys.readouterr().out

    assert exit_code == 0
    assert "Validation passed" in output


def test_cli_returns_failure_for_invalid_file(tmp_path, capsys):
    data = pd.DataFrame(
        {
            "name": ["Alice", "Bob"],
        }
    )

    file_path = tmp_path / "invalid.csv"
    data.to_csv(file_path, index=False)

    exit_code = main(
        [
            str(file_path),
            "--required",
            "name",
            "email",
        ]
    )

    output = capsys.readouterr().out

    assert exit_code == 1
    assert "Missing columns: email" in output

def test_cli_uses_yaml_config(tmp_path, capsys):
    data = pd.DataFrame(
        {
            "id": [101, 101],
            "name": ["Alice", "Alice"],
            "employment_type": ["full-time", "temporary"],
        }
    )

    data_file = tmp_path / "jobs.csv"
    data.to_csv(data_file, index=False)

    config_file = tmp_path / "rules.yaml"
    config_file.write_text(
        """
required_columns:
  - id
  - name
  - employment_type

expected_columns:
  - id
  - name
  - employment_type

duplicate_subset:
  - id

allowed_values:
  employment_type:
    - full-time
    - part-time
    - contract
""",
        encoding="utf-8",
    )

    exit_code = main(
        [
            str(data_file),
            "--config",
            str(config_file),
        ]
    )

    output = capsys.readouterr().out

    assert exit_code == 1
    assert "Duplicate rows: 3" in output
    assert "Invalid values in 'employment_type': rows 3" in output


def test_cli_config_passes_valid_data(tmp_path, capsys):
    data = pd.DataFrame(
        {
            "id": [101, 102],
            "name": ["Alice", "Bob"],
            "employment_type": ["full-time", "contract"],
        }
    )

    data_file = tmp_path / "jobs.csv"
    data.to_csv(data_file, index=False)

    config_file = tmp_path / "rules.yaml"
    config_file.write_text(
        """
required_columns:
  - id
  - name
  - employment_type

expected_columns:
  - id
  - name
  - employment_type

allowed_values:
  employment_type:
    - full-time
    - part-time
    - contract
""",
        encoding="utf-8",
    )

    exit_code = main(
        [
            str(data_file),
            "--config",
            str(config_file),
        ]
    )

    output = capsys.readouterr().out

    assert exit_code == 0
    assert "Validation passed" in output

def test_cli_exports_csv_report(tmp_path):
    data = pd.DataFrame(
        {
            "name": ["Alice", "Bob"],
        }
    )

    data_file = tmp_path / "data.csv"
    data.to_csv(data_file, index=False)

    output_file = tmp_path / "report.csv"

    exit_code = main(
        [
            str(data_file),
            "--required",
            "name",
            "email",
            "--output",
            str(output_file),
        ]
    )

    report = pd.read_csv(output_file)

    assert exit_code == 1
    assert output_file.exists()
    assert "missing_column" in report["issue_type"].values
    assert "email" in report["column"].values

from openpyxl import load_workbook


def test_cli_exports_excel_report(tmp_path):
    data = pd.DataFrame(
        {
            "name": ["Alice", "Bob"],
        }
    )

    data_file = tmp_path / "data.csv"
    data.to_csv(data_file, index=False)

    output_file = tmp_path / "report.xlsx"

    exit_code = main(
        [
            str(data_file),
            "--required",
            "name",
            "email",
            "--output",
            str(output_file),
        ]
    )

    workbook = load_workbook(output_file)

    assert exit_code == 1
    assert output_file.exists()
    assert "Summary" in workbook.sheetnames
    assert "Issues" in workbook.sheetnames

def test_cli_handles_missing_input_file(tmp_path, capsys):
    missing_file = tmp_path / "missing.csv"

    exit_code = main([str(missing_file)])

    output = capsys.readouterr().out

    assert exit_code == 2
    assert "Error:" in output
    assert "File not found" in output


def test_cli_handles_missing_config_file(tmp_path, capsys):
    data = pd.DataFrame(
        {
            "name": ["Alice"],
        }
    )

    data_file = tmp_path / "data.csv"
    data.to_csv(data_file, index=False)

    missing_config = tmp_path / "missing.yaml"

    exit_code = main(
        [
            str(data_file),
            "--config",
            str(missing_config),
        ]
    )

    output = capsys.readouterr().out

    assert exit_code == 2
    assert "Error:" in output
    assert "Configuration file not found" in output


def test_cli_handles_invalid_config(tmp_path, capsys):
    data = pd.DataFrame(
        {
            "name": ["Alice"],
        }
    )

    data_file = tmp_path / "data.csv"
    data.to_csv(data_file, index=False)

    config_file = tmp_path / "rules.yaml"
    config_file.write_text(
        """
required_columns: name
""",
        encoding="utf-8",
    )

    exit_code = main(
        [
            str(data_file),
            "--config",
            str(config_file),
        ]
    )

    output = capsys.readouterr().out

    assert exit_code == 2
    assert "Error:" in output
    assert "required_columns must be a list" in output