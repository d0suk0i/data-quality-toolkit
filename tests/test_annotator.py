from openpyxl import Workbook, load_workbook

from src.data_quality_toolkit.annotator import annotate_workbook
from src.data_quality_toolkit.validator import ValidationResult


def test_annotate_workbook_highlights_validation_issues(tmp_path):
    source_path = tmp_path / "source.xlsx"
    output_path = tmp_path / "reviewed.xlsx"

    workbook = Workbook()
    sheet = workbook.active

    sheet.append(
        [
            "id",
            "name",
            "email",
            "employment_type",
            "extra",
        ]
    )

    sheet.append(
        [
            101,
            "Alice",
            None,
            "full-time",
            "A",
        ]
    )

    sheet.append(
        [
            102,
            "Bob",
            "bob@example.com",
            "temporary",
            "B",
        ]
    )

    sheet.append(
        [
            102,
            "Bob",
            "bob@example.com",
            "contract",
            "C",
        ]
    )

    workbook.save(source_path)

    result = ValidationResult(
        unexpected_columns=["extra"],
        missing_values={
            "email": [0],
        },
        invalid_values={
            "employment_type": [1],
        },
        duplicate_rows=[2],
    )

    annotate_workbook(
        input_path=source_path,
        result=result,
        output_path=output_path,
    )

    reviewed = load_workbook(output_path)
    reviewed_sheet = reviewed.active

    assert output_path.exists()

    # Original data is preserved.
    assert reviewed_sheet["A2"].value == 101
    assert reviewed_sheet["B3"].value == "Bob"

    # Missing email.
    assert reviewed_sheet["C2"].fill.fill_type == "solid"

    # Invalid employment type.
    assert reviewed_sheet["D3"].fill.fill_type == "solid"

    # Duplicate row.
    assert reviewed_sheet["A4"].fill.fill_type == "solid"
    assert reviewed_sheet["E4"].fill.fill_type == "solid"

    # Unexpected column header.
    assert reviewed_sheet["E1"].fill.fill_type == "solid"